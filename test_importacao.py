import unittest
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
from unittest.mock import patch

from docx import Document

import importacao as imp
from ocr import (
    ErroOCR,
    ResultadoOCR,
)


class TestImportacao(unittest.TestCase):

    def test_txt_preserva_acentos(self):
        dados = (
            "Solicitação de material."
            .encode(
                "utf-8-sig"
            )
        )

        self.assertEqual(
            imp.extrair_texto(
                "teste.TXT",
                dados,
            ),
            "Solicitação de material.",
        )

    def test_docx_le_paragrafos_e_tabela_em_ordem(self):
        documento = Document()

        documento.add_paragraph(
            "Antes da tabela."
        )

        tabela = documento.add_table(
            rows=1,
            cols=2,
        )

        tabela.cell(
            0,
            0,
        ).text = "Setor"

        tabela.cell(
            0,
            1,
        ).text = "RH"

        documento.add_paragraph(
            "Depois da tabela."
        )

        memoria = BytesIO()

        documento.save(
            memoria
        )

        texto = imp.extrair_texto(
            "teste.docx",
            memoria.getvalue(),
        )

        self.assertEqual(
            texto,
            (
                "Antes da tabela.\n\n"
                "Setor | RH\n\n"
                "Depois da tabela."
            ),
        )

    def test_pdf_identifica_pagina_com_texto_nativo(self):
        leitor = SimpleNamespace(
            is_encrypted=False,
            pages=[
                SimpleNamespace(
                    extract_text=lambda: (
                        "Material fictício."
                    )
                )
            ],
        )

        with patch(
            "importacao.PdfReader",
            return_value=leitor,
        ):
            with patch(
                "importacao.reconhecer_pagina_pdf"
            ) as ocr_pdf:
                texto = imp.extrair_texto(
                    "teste.pdf",
                    b"simulado",
                )

        self.assertEqual(
            texto,
            (
                "[Página 1]\n"
                "Material fictício."
            ),
        )

        ocr_pdf.assert_not_called()

    def test_pdf_digitalizado_usa_ocr(self):
        leitor = SimpleNamespace(
            is_encrypted=False,
            pages=[
                SimpleNamespace(
                    extract_text=lambda: ""
                )
            ],
        )

        resultado = ResultadoOCR(
            texto=(
                "Procedimento digitalizado."
            ),
            confianca_media=0.91,
            linhas_reconhecidas=1,
        )

        with patch(
            "importacao.PdfReader",
            return_value=leitor,
        ):
            with patch(
                "importacao.reconhecer_pagina_pdf",
                return_value=resultado,
            ) as ocr_pdf:
                texto = imp.extrair_texto(
                    "digitalizado.pdf",
                    b"pdf-simulado",
                )

        self.assertEqual(
            texto,
            (
                "[Página 1]\n"
                "[Extração: OCR]\n"
                "Procedimento digitalizado."
            ),
        )

        ocr_pdf.assert_called_once_with(
            b"pdf-simulado",
            0,
        )

    def test_pdf_misto_usa_ocr_apenas_na_pagina_sem_texto(self):
        leitor = SimpleNamespace(
            is_encrypted=False,
            pages=[
                SimpleNamespace(
                    extract_text=lambda: (
                        "Página com texto normal."
                    )
                ),
                SimpleNamespace(
                    extract_text=lambda: ""
                ),
                SimpleNamespace(
                    extract_text=lambda: (
                        "Outra página normal."
                    )
                ),
            ],
        )

        resultado = ResultadoOCR(
            texto=(
                "Página digitalizada."
            ),
            confianca_media=0.88,
            linhas_reconhecidas=1,
        )

        with patch(
            "importacao.PdfReader",
            return_value=leitor,
        ):
            with patch(
                "importacao.reconhecer_pagina_pdf",
                return_value=resultado,
            ) as ocr_pdf:
                texto = imp.extrair_texto(
                    "misto.pdf",
                    b"pdf-misto",
                )

        self.assertEqual(
            texto,
            (
                "[Página 1]\n"
                "Página com texto normal.\n\n"
                "[Página 2]\n"
                "[Extração: OCR]\n"
                "Página digitalizada.\n\n"
                "[Página 3]\n"
                "Outra página normal."
            ),
        )

        ocr_pdf.assert_called_once_with(
            b"pdf-misto",
            1,
        )

    def test_pdf_preserva_numero_real_da_pagina_com_ocr(self):
        leitor = SimpleNamespace(
            is_encrypted=False,
            pages=[
                SimpleNamespace(
                    extract_text=lambda: (
                        "Primeira."
                    )
                ),
                SimpleNamespace(
                    extract_text=lambda: (
                        "Segunda."
                    )
                ),
                SimpleNamespace(
                    extract_text=lambda: ""
                ),
            ],
        )

        resultado = ResultadoOCR(
            texto="Terceira via OCR.",
            confianca_media=0.9,
            linhas_reconhecidas=1,
        )

        with patch(
            "importacao.PdfReader",
            return_value=leitor,
        ):
            with patch(
                "importacao.reconhecer_pagina_pdf",
                return_value=resultado,
            ):
                texto = imp.extrair_texto(
                    "teste.pdf",
                    b"pdf",
                )

        self.assertIn(
            (
                "[Página 3]\n"
                "[Extração: OCR]\n"
                "Terceira via OCR."
            ),
            texto,
        )

    def test_pdf_falha_se_ocr_nao_reconhecer_pagina(self):
        leitor = SimpleNamespace(
            is_encrypted=False,
            pages=[
                SimpleNamespace(
                    extract_text=lambda: ""
                )
            ],
        )

        with patch(
            "importacao.PdfReader",
            return_value=leitor,
        ):
            with patch(
                "importacao.reconhecer_pagina_pdf",
                side_effect=ErroOCR(
                    "Não encontrei texto legível na imagem."
                ),
            ):
                with self.assertRaisesRegex(
                    imp.ErroImportacao,
                    "Página 1",
                ):
                    imp.extrair_texto(
                        "teste.pdf",
                        b"pdf",
                    )

    def test_pdf_protegido_continua_rejeitado(self):
        leitor = SimpleNamespace(
            is_encrypted=True,
            pages=[],
        )

        with patch(
            "importacao.PdfReader",
            return_value=leitor,
        ):
            with self.assertRaisesRegex(
                imp.ErroImportacao,
                "protegido",
            ):
                imp.extrair_texto(
                    "teste.pdf",
                    b"pdf",
                )

    def test_pdf_limita_quantidade_de_paginas(self):
        paginas = [
            SimpleNamespace(
                extract_text=lambda: (
                    "Texto."
                )
            )
            for _ in range(
                imp.MAX_PAGINAS + 1
            )
        ]

        leitor = SimpleNamespace(
            is_encrypted=False,
            pages=paginas,
        )

        with patch(
            "importacao.PdfReader",
            return_value=leitor,
        ):
            with self.assertRaisesRegex(
                imp.ErroImportacao,
                "até 30 páginas",
            ):
                imp.extrair_texto(
                    "teste.pdf",
                    b"pdf",
                )

    def test_imagem_png_usa_ocr(self):
        resultado = ResultadoOCR(
            texto=(
                "Texto reconhecido "
                "na imagem."
            ),
            confianca_media=0.94,
            linhas_reconhecidas=1,
        )

        with patch(
            "importacao.reconhecer_imagem",
            return_value=resultado,
        ) as reconhecer:
            texto = imp.extrair_texto(
                "captura.PNG",
                b"imagem-ficticia",
            )

        self.assertEqual(
            texto,
            (
                "[Extração: OCR]\n"
                "Texto reconhecido na imagem."
            ),
        )

        reconhecer.assert_called_once_with(
            b"imagem-ficticia"
        )

    def test_formatos_de_imagem_sao_aceitos(self):
        resultado = ResultadoOCR(
            texto="Texto.",
            confianca_media=0.9,
            linhas_reconhecidas=1,
        )

        formatos = [
            "imagem.png",
            "imagem.jpg",
            "imagem.jpeg",
            "imagem.bmp",
            "imagem.tif",
            "imagem.tiff",
            "imagem.webp",
        ]

        with patch(
            "importacao.reconhecer_imagem",
            return_value=resultado,
        ):
            for nome in formatos:
                with self.subTest(
                    nome=nome
                ):
                    self.assertEqual(
                        imp.extrair_texto(
                            nome,
                            b"imagem",
                        ),
                        (
                            "[Extração: OCR]\n"
                            "Texto."
                        ),
                    )

    def test_imagem_sem_texto_legivel_e_rejeitada(self):
        with patch(
            "importacao.reconhecer_imagem",
            side_effect=ErroOCR(
                "Não encontrei texto legível na imagem."
            ),
        ):
            with self.assertRaisesRegex(
                imp.ErroImportacao,
                "OCR",
            ):
                imp.extrair_texto(
                    "imagem.jpg",
                    b"imagem",
                )

    def test_rejeita_formatos_conteudos_e_tamanho_invalidos(self):
        casos = [
            (
                "a.exe",
                b"abc",
            ),
            (
                "a.txt",
                b"",
            ),
            (
                "a.txt",
                b"   ",
            ),
            (
                "a.txt",
                b"\xff",
            ),
        ]

        for nome, dados in casos:
            with self.subTest(
                nome=nome,
                dados=dados,
            ):
                with self.assertRaises(
                    imp.ErroImportacao
                ):
                    imp.extrair_texto(
                        nome,
                        dados,
                    )

        with patch(
            "importacao.MAX_BYTES",
            2,
        ):
            with self.assertRaises(
                imp.ErroImportacao
            ):
                imp.extrair_texto(
                    "a.txt",
                    b"abc",
                )

    def test_rejeita_doc_antigo(self):
        with self.assertRaises(
            imp.ErroImportacao
        ):
            imp.extrair_texto(
                "antigo.doc",
                b"conteudo",
            )

    def test_rejeita_texto_extraido_acima_do_limite(self):
        with patch(
            "importacao.MAX_TEXTO",
            3,
        ):
            with self.assertRaises(
                imp.ErroImportacao
            ):
                imp.extrair_texto(
                    "teste.txt",
                    b"abcd",
                )

    def test_salva_apenas_na_pasta_de_destino(self):
        with TemporaryDirectory() as temporaria:
            pasta = (
                Path(temporaria)
                / "documentos"
            )

            destino, criado = (
                imp.salvar_texto(
                    pasta,
                    "../../fora.txt",
                    "Conteúdo fictício.",
                )
            )

            self.assertTrue(
                criado
            )

            self.assertEqual(
                destino.parent,
                pasta,
            )

            self.assertTrue(
                destino.name.startswith(
                    "importado_"
                )
            )

            self.assertIn(
                "Conteúdo fictício.",
                destino.read_text(
                    encoding="utf-8"
                ),
            )

    def test_salvamento_preserva_nome_original_da_imagem(self):
        with TemporaryDirectory() as temporaria:
            pasta = Path(
                temporaria
            )

            destino, criado = (
                imp.salvar_texto(
                    pasta,
                    "captura.png",
                    (
                        "[Extração: OCR]\n"
                        "Texto reconhecido."
                    ),
                )
            )

            self.assertTrue(
                criado
            )

            conteudo = (
                destino.read_text(
                    encoding="utf-8"
                )
            )

            self.assertIn(
                (
                    "Arquivo de origem: "
                    "captura.png"
                ),
                conteudo,
            )

            self.assertIn(
                "[Extração: OCR]",
                conteudo,
            )

    def test_reenvio_nao_duplica_nem_sobrescreve(self):
        with TemporaryDirectory() as temporaria:
            pasta = Path(
                temporaria
            )

            primeiro, _ = (
                imp.salvar_texto(
                    pasta,
                    "teste.txt",
                    "Conteúdo.",
                )
            )

            segundo, criado = (
                imp.salvar_texto(
                    pasta,
                    "teste.txt",
                    "Conteúdo.",
                )
            )

            self.assertEqual(
                primeiro,
                segundo,
            )

            self.assertFalse(
                criado
            )

            self.assertEqual(
                len(
                    list(
                        pasta.glob(
                            "*.txt"
                        )
                    )
                ),
                1,
            )

    def test_texto_alterado_cria_outra_copia(self):
        with TemporaryDirectory() as temporaria:
            pasta = Path(
                temporaria
            )

            primeiro, _ = (
                imp.salvar_texto(
                    pasta,
                    "teste.txt",
                    "Versão inicial.",
                )
            )

            segundo, criado = (
                imp.salvar_texto(
                    pasta,
                    "teste.txt",
                    "Versão nova.",
                )
            )

            self.assertTrue(
                criado
            )

            self.assertNotEqual(
                primeiro,
                segundo,
            )

            self.assertIn(
                "Versão inicial.",
                primeiro.read_text(
                    encoding="utf-8"
                ),
            )


if __name__ == "__main__":
    unittest.main()